"""
Document Generator for HavenSec Arena College Deliverables.
Generates HavenSec_Arena_Final_Report.docx and HavenSec_Arena_Final_Report.pdf
directly from HavenSec_Arena_Final_Report.md.
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

BASE_DIR = os.path.dirname(__file__)
MD_FILE = os.path.join(BASE_DIR, "HavenSec_Arena_Final_Report.md")
DOCX_FILE = os.path.join(BASE_DIR, "HavenSec_Arena_Final_Report.docx")
PDF_FILE = os.path.join(BASE_DIR, "HavenSec_Arena_Final_Report.pdf")

def generate_docx():
    print("Generating DOCX...")
    doc = Document()

    # Set 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    with open(MD_FILE, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle Tables
        if "|" in stripped and ("---" in stripped or stripped.startswith("|")):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            # Table ended, process it
            process_docx_table(doc, table_lines)
            in_table = False
            table_lines = []

        if not stripped:
            continue

        if stripped.startswith("# "):
            h = doc.add_heading(stripped[2:], level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            add_formatted_text(p, stripped[2:])
        elif stripped.startswith("```"):
            continue
        else:
            p = doc.add_paragraph()
            add_formatted_text(p, stripped)

    if in_table and table_lines:
        process_docx_table(doc, table_lines)

    doc.save(DOCX_FILE)
    print(f"Saved: {DOCX_FILE}")

def add_formatted_text(paragraph, text):
    # Basic markdown bold (**text**) handling
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)

def process_docx_table(doc, table_lines):
    rows = []
    for l in table_lines:
        if "---" in l:
            continue
        cells = [c.strip() for c in l.split("|")[1:-1]]
        if cells:
            rows.append(cells)

    if not rows:
        return

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    for r_idx, row in enumerate(rows):
        for c_idx, cell in enumerate(row):
            if c_idx < len(table.rows[r_idx].cells):
                cell_p = table.cell(r_idx, c_idx).paragraphs[0]
                cell_p.text = cell
                if r_idx == 0:
                    for run in cell_p.runs:
                        run.bold = True

def generate_pdf():
    print("Generating PDF...")
    doc = SimpleDocTemplate(
        PDF_FILE,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=14
    )
    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#1d4ed8'),
        spaceBefore=14,
        spaceAfter=8
    )
    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#0f172a'),
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )
    bullet_style = ParagraphStyle(
        'Bullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        leftIndent=15,
        spaceAfter=4
    )
    code_style = ParagraphStyle(
        'Code',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=6
    )

    with open(MD_FILE, "r", encoding="utf-8") as f:
        lines = f.readlines()

    story = []
    in_table = False
    table_lines = []

    for line in lines:
        stripped = line.strip()

        # Table handling
        if "|" in stripped and ("---" in stripped or stripped.startswith("|")):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            story.append(process_pdf_table(table_lines, body_style))
            story.append(Spacer(1, 8))
            in_table = False
            table_lines = []

        if not stripped:
            continue

        # Format markdown bold and handle HTML tags for ReportLab
        formatted = re.sub(r'<br>', r'<br/>', stripped)
        formatted = re.sub(r'&nbsp;', r' ', formatted)
        formatted = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', formatted)
        formatted = re.sub(r'`(.*?)`', r'<font face="Courier">\1</font>', formatted)

        # Escape isolated ampersands that are not XML entities
        formatted = re.sub(r'&(?!(amp|lt|gt|quot|apos|bull);)', '&amp;', formatted)

        if stripped.startswith("# "):
            story.append(Paragraph(formatted[2:], title_style))
        elif stripped.startswith("## "):
            story.append(Paragraph(formatted[3:], h1_style))
        elif stripped.startswith("### "):
            story.append(Paragraph(formatted[4:], h2_style))
        elif stripped.startswith("- "):
            story.append(Paragraph(f"&bull; {formatted[2:]}", bullet_style))
        elif stripped.startswith("```"):
            continue
        else:
            story.append(Paragraph(formatted, body_style))

    if in_table and table_lines:
        story.append(process_pdf_table(table_lines, body_style))

    doc.build(story)
    print(f"Saved: {PDF_FILE}")

def process_pdf_table(table_lines, style):
    rows = []
    for l in table_lines:
        if "---" in l:
            continue
        cells = [c.strip() for c in l.split("|")[1:-1]]
        if cells:
            formatted_cells = [Paragraph(re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', c), style) for c in cells]
            rows.append(formatted_cells)

    if not rows:
        return Spacer(1, 1)

    t = Table(rows, colWidths=[500 / len(rows[0])] * len(rows[0]))
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
    ]))
    return t

if __name__ == "__main__":
    generate_docx()
    generate_pdf()
